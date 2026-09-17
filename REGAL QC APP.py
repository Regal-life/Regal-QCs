from datetime import datetime
import io
import json
import sqlite3

from fpdf import FPDF
import openpyxl
from openpyxl.drawing.image import Image as OpenPyxlImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from PIL import Image
import streamlit as st

# Optional Barcode Reader library
try:
    from pyzbar.pyzbar import decode as decode_barcode

    HAS_PYZBAR = True
except ImportError:
    HAS_PYZBAR = False

# Page Configuration
st.set_page_config(
    page_title="Regal QC Manager",
    page_icon="📋",
    layout="centered",
    initial_sidebar_state="collapsed",
)

# Custom CSS Styling
st.markdown(
    """
    <style>
    .main .block-container { padding-top: 1.5rem; padding-bottom: 3rem; max-width: 800px; }
    .app-header { background: linear-gradient(135deg, #1F497D 0%, #112948 100%); padding: 1.2rem; border-radius: 12px; color: white; text-align: center; margin-bottom: 1.2rem; box-shadow: 0 4px 12px rgba(0, 0, 0, 0.1); }
    .app-header h1 { margin: 0; font-size: 1.7rem; font-weight: 700; color: #ffffff; }
    .app-header p { margin: 0.2rem 0 0 0; font-size: 0.9rem; opacity: 0.85; }
    .metric-card { background-color: #ffffff; border: 1px solid #E2E8F0; border-radius: 10px; padding: 0.8rem; text-align: center; }
    .metric-value { font-size: 1.5rem; font-weight: 700; color: #1F497D; }
    .metric-label { font-size: 0.75rem; color: #64748B; text-transform: uppercase; }
    .badge-resolved { background-color: #DEF7EC; color: #03543F; padding: 3px 8px; border-radius: 12px; font-size: 0.75rem; font-weight: 600; }
    .badge-open { background-color: #FDE8E8; color: #9B1C1C; padding: 3px 8px; border-radius: 12px; font-size: 0.75rem; font-weight: 600; }
    </style>
""",
    unsafe_allow_html=True,
)


# --- DATABASE SETUP ---
def init_db():
    conn = sqlite3.connect("qc_logs.db")
    cursor = conn.cursor()
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS qc_entries (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            date_opened TEXT,
            product_code TEXT,
            quantity TEXT,
            time_taken TEXT,
            supplier TEXT,
            resolved TEXT,
            date_completed TEXT,
            product_desc TEXT,
            issue_desc TEXT,
            corrective_action TEXT,
            photo_data BLOB,
            photos_json TEXT,
            timestamp DATETIME DEFAULT CURRENT_TIMESTAMP
        )
    """)
    conn.commit()
    conn.close()


init_db()


def add_qc_entry(entry):
    conn = sqlite3.connect("qc_logs.db")
    cursor = conn.cursor()
    cursor.execute(
        """
        INSERT INTO qc_entries (
            date_opened, product_code, quantity, time_taken, supplier, 
            resolved, date_completed, product_desc, issue_desc, corrective_action, photo_data, photos_json
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
    """,
        (
            entry["date_opened"],
            entry["product_code"],
            entry["quantity"],
            entry["time_taken"],
            entry["supplier"],
            entry["resolved"],
            entry["date_completed"],
            entry["product_desc"],
            entry["issue_desc"],
            entry["corrective_action"],
            entry["photo_data"],
            entry["photos_json"],
        ),
    )
    conn.commit()
    conn.close()


def get_all_entries():
    conn = sqlite3.connect("qc_logs.db")
    cursor = conn.cursor()
    cursor.execute(
        "SELECT id, date_opened, product_code, quantity, photo_data, product_desc, issue_desc, corrective_action, resolved, date_completed, time_taken, supplier, photos_json FROM qc_entries ORDER BY id DESC"
    )
    rows = cursor.fetchall()
    conn.close()
    return rows


def delete_entry(entry_id):
    conn = sqlite3.connect("qc_logs.db")
    cursor = conn.cursor()
    cursor.execute("DELETE FROM qc_entries WHERE id = ?", (entry_id,))
    conn.commit()
    conn.close()


# --- PDF CLEARANCE CERTIFICATE GENERATOR ---
def generate_pdf_certificate(rec):
    (
        rec_id,
        d_open,
        p_code,
        qty,
        photo,
        p_desc,
        issue,
        action,
        res,
        d_comp,
        t_taken,
        supp,
        p_json,
    ) = rec

    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 16)
    pdf.cell(0, 10, "REGAL DISTRIBUTORS SA", ln=True, align="C")
    pdf.set_font("Helvetica", "", 12)
    pdf.cell(0, 8, "QUALITY CONTROL CLEARANCE CERTIFICATE", ln=True, align="C")
    pdf.line(10, 30, 200, 30)
    pdf.ln(10)

    pdf.set_font("Helvetica", "B", 11)
    pdf.cell(50, 8, "Case Ref Number:", border=0)
    pdf.set_font("Helvetica", "", 11)
    pdf.cell(0, 8, f"#{rec_id}", ln=True)

    pdf.set_font("Helvetica", "B", 11)
    pdf.cell(50, 8, "Product Code:", border=0)
    pdf.set_font("Helvetica", "", 11)
    pdf.cell(0, 8, f"{p_code}", ln=True)

    pdf.set_font("Helvetica", "B", 11)
    pdf.cell(50, 8, "Supplier / Location:", border=0)
    pdf.set_font("Helvetica", "", 11)
    pdf.cell(0, 8, f"{supp}", ln=True)

    pdf.set_font("Helvetica", "B", 11)
    pdf.cell(50, 8, "Status:", border=0)
    pdf.set_font("Helvetica", "", 11)
    pdf.cell(
        0,
        8,
        f"{'RESOLVED / APPROVED FOR SALE' if res == 'Y' else 'OPEN / BLOCKED'}",
        ln=True,
    )

    pdf.ln(5)
    pdf.set_font("Helvetica", "B", 11)
    pdf.cell(0, 8, "Product Description:", ln=True)
    pdf.set_font("Helvetica", "", 10)
    pdf.multi_cell(0, 6, p_desc if p_desc else "N/A")

    pdf.ln(3)
    pdf.set_font("Helvetica", "B", 11)
    pdf.cell(0, 8, "Issue / Defect Description:", ln=True)
    pdf.set_font("Helvetica", "", 10)
    pdf.multi_cell(0, 6, issue if issue else "N/A")

    pdf.ln(3)
    pdf.set_font("Helvetica", "B", 11)
    pdf.cell(0, 8, "Corrective Action / Outcome:", ln=True)
    pdf.set_font("Helvetica", "", 10)
    pdf.multi_cell(0, 6, action if action else "N/A")

    pdf.ln(15)
    pdf.set_font("Helvetica", "I", 9)
    pdf.cell(
        0,
        6,
        f"Generated automatically on {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}",
        ln=True,
        align="R",
    )

    return bytes(pdf.output())


# --- APP HEADER ---
st.markdown(
    """
    <div class="app-header">
        <h1>📋 Regal Quality Control</h1>
        <p>Mobile Inspection & Report Management</p>
    </div>
""",
    unsafe_allow_html=True,
)

tab1, tab2, tab3 = st.tabs(
    ["➕ New Case", "📜 History Logs", "📊 Generate Report"]
)

# --- TAB 1: NEW INCIDENT ---
with tab1:
    st.caption("Record new inspection details or scan item packaging.")

    # 1. BARCODE / QR SCANNER
    with st.expander("📷 Scan Packaging Barcode (Auto-Fill Code)"):
        barcode_cam = st.camera_input("Point camera at barcode", key="barcode_scan")
        if barcode_cam and HAS_PYZBAR:
            img = Image.open(barcode_cam)
            decoded = decode_barcode(img)
            if decoded:
                scanned_code = decoded[0].data.decode("utf-8")
                st.session_state["scanned_product_code"] = scanned_code
                st.success(f"Barcode Detected: **{scanned_code}**")
            else:
                st.warning("No barcode detected. Try re-positioning.")

    # 2. ENTRY FORM
    scanned_val = st.session_state.get("scanned_product_code", "")

    with st.form("qc_entry_form", clear_on_submit=True):
        col1, col2 = st.columns(2)
        with col1:
            date_opened = st.date_input("Date Opened", value=datetime.today())
            product_code = st.text_input(
                "Product Code",
                value=scanned_val,
                placeholder="e.g. IC118-4-1",
            )
            quantity = st.text_input("Quantity", value="1")
            time_taken = st.text_input("Time Taken", value="1 Hour")

        with col2:
            supplier = st.text_input(
                "Supplier / Location", placeholder="e.g. Accentronix CC"
            )
            resolved = st.selectbox(
                "Status",
                ["Y", "N"],
                format_func=lambda x: "Resolved (Y)" if x == "Y" else "Open (N)",
            )
            date_completed = st.date_input("Date Completed", value=datetime.today())

        st.divider()
        product_desc = st.text_area(
            "Product Description", placeholder="Enter item details..."
        )
        issue_desc = st.text_area(
            "Issue Description", placeholder="Describe defect or query..."
        )
        corrective_action = st.text_area(
            "Corrective Action", placeholder="Action taken..."
        )

        st.divider()
        st.markdown("📷 **Multi-Photo Inspection Attachment**")
        uploaded_photos = st.file_uploader(
            "Upload or capture defect photos",
            type=["png", "jpg", "jpeg"],
            accept_multiple_files=True,
        )

        submitted = st.form_submit_button(
            "💾 Save Case Entry", use_container_width=True, type="primary"
        )

        if submitted:
            if not product_code:
                st.error("Please enter a Product Code.")
            else:
                encoded_photos = []
                if uploaded_photos:
                    for p in uploaded_photos:
                        encoded_photos.append(p.getvalue().hex())

                first_photo = (
                    uploaded_photos[0].getvalue() if uploaded_photos else None
                )

                entry = {
                    "date_opened": date_opened.strftime("%Y-%m-%d"),
                    "product_code": product_code,
                    "quantity": quantity,
                    "time_taken": time_taken,
                    "supplier": supplier,
                    "resolved": resolved,
                    "date_completed": date_completed.strftime("%Y-%m-%d")
                    if resolved == "Y"
                    else "",
                    "product_desc": product_desc,
                    "issue_desc": issue_desc,
                    "corrective_action": corrective_action,
                    "photo_data": first_photo,
                    "photos_json": json.dumps(encoded_photos),
                }
                add_qc_entry(entry)
                st.toast("Case saved to database!", icon="✅")
                st.success(f"Case logged for product {product_code}.")


# --- TAB 2: HISTORY & LOGS ---
with tab2:
    records = get_all_entries()

    total_cases = len(records)
    resolved_cases = sum(1 for r in records if r[8] == "Y")
    open_cases = total_cases - resolved_cases

    m1, m2, m3 = st.columns(3)
    with m1:
        st.markdown(
            f'<div class="metric-card"><div class="metric-value">{total_cases}</div><div class="metric-label">Total Cases</div></div>',
            unsafe_allow_html=True,
        )
    with m2:
        st.markdown(
            f'<div class="metric-card"><div class="metric-value" style="color: #059669;">{resolved_cases}</div><div class="metric-label">Resolved</div></div>',
            unsafe_allow_html=True,
        )
    with m3:
        st.markdown(
            f'<div class="metric-card"><div class="metric-value" style="color: #DC2626;">{open_cases}</div><div class="metric-label">Open</div></div>',
            unsafe_allow_html=True,
        )

    st.write("")

    if not records:
        st.info("No recorded logs in the database.")
    else:
        for rec in records:
            (
                rec_id,
                d_open,
                p_code,
                qty,
                photo,
                p_desc,
                issue,
                action,
                res,
                d_comp,
                t_taken,
                supp,
                p_json,
            ) = rec
            badge_html = (
                '<span class="badge-resolved">Resolved</span>'
                if res == "Y"
                else '<span class="badge-open">Open</span>'
            )

            with st.expander(f"Case #{rec_id} — {p_code}"):
                st.markdown(
                    f"**Supplier:** {supp} | **Status:** {badge_html}",
                    unsafe_allow_html=True,
                )
                st.write("")

                c1, c2 = st.columns([2, 1])
                with c1:
                    st.markdown(f"**Date Opened:** {d_open}")
                    st.markdown(f"**Quantity:** {qty} | **Time:** {t_taken}")
                    if p_desc:
                        st.markdown(f"**Description:** {p_desc}")
                    if issue:
                        st.markdown(f"**Issue:** {issue}")
                    if action:
                        st.markdown(f"**Action:** {action}")
                with c2:
                    if p_json:
                        try:
                            hex_list = json.loads(p_json)
                            for hex_img in hex_list[:3]:
                                st.image(
                                    bytes.fromhex(hex_img),
                                    use_container_width=True,
                                )
                        except Exception:
                            if photo:
                                st.image(photo, use_container_width=True)
                    elif photo:
                        st.image(photo, use_container_width=True)

                st.divider()
                col_btn1, col_btn2 = st.columns(2)
                with col_btn1:
                    pdf_bytes = generate_pdf_certificate(rec)
                    st.download_button(
                        label="📄 Export PDF Clearance",
                        data=pdf_bytes,
                        file_name=f"QC_Certificate_Case_{rec_id}.pdf",
                        mime="application/pdf",
                        use_container_width=True,
                    )
                with col_btn2:
                    if st.button(
                        f"🗑️ Delete Entry",
                        key=f"del_{rec_id}",
                        type="secondary",
                        use_container_width=True,
                    ):
                        delete_entry(rec_id)
                        st.toast(f"Deleted Case #{rec_id}", icon="🗑️")
                        st.rerun()


# --- TAB 3: EXCEL GENERATOR ENGINE ---
def build_excel_report(month_year_str, records):
    wb = openpyxl.Workbook()

    fill_header = PatternFill(
        start_color="1F497D", end_color="1F497D", fill_type="solid"
    )
    font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    font_title = Font(name="Calibri", size=16, bold=True)
    border_grid = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    ws_summary = wb.active
    ws_summary.title = "QC Presentation"
    ws_summary.views.sheetView[0].showGridLines = True
    ws_summary["A1"] = f"{month_year_str.upper()} QUALITY CONTROL — SUMMARY"
    ws_summary["A1"].font = font_title

    ws_reg = wb.create_sheet(title=f"QC {month_year_str}")
    ws_reg.views.sheetView[0].showGridLines = True
    ws_reg["A1"] = f"QUALITY CONTROL REGISTER — {month_year_str.upper()}"
    ws_reg["A1"].font = font_title

    headers = [
        "No.",
        "Date Opened/Checked",
        "Product Code",
        "Quantity",
        "Picture of Product",
        "Product Description",
        "Background / Issue Description",
        "Corrective Action",
        "Resolved (Y/N)",
        "Date Completed",
        "Time Taken",
        "Supplier",
    ]

    for col_idx, header in enumerate(headers, 1):
        cell = ws_reg.cell(row=3, column=col_idx, value=header)
        cell.fill = fill_header
        cell.font = font_header
        cell.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )

    for row_idx, rec in enumerate(records, 4):
        (
            rec_id,
            d_open,
            p_code,
            qty,
            photo,
            p_desc,
            issue,
            action,
            res,
            d_comp,
            t_taken,
            supp,
            p_json,
        ) = rec
        ws_reg.row_dimensions[row_idx].height = 65 if photo else 20

        row_vals = [
            rec_id,
            d_open,
            p_code,
            qty,
            "",
            p_desc,
            issue,
            action,
            res,
            d_comp,
            t_taken,
            supp,
        ]

        for col_idx, val in enumerate(row_vals, 1):
            cell = ws_reg.cell(row=row_idx, column=col_idx, value=val)
            cell.border = border_grid
            if col_idx in [1, 4, 9]:
                cell.alignment = Alignment(horizontal="center")

        if photo:
            img_bytes = io.BytesIO(photo)
            img = Image.open(img_bytes)
            img.thumbnail((80, 80))
            img_out = io.BytesIO()
            img.save(img_out, format="PNG")
            img_out.seek(0)

            xl_img = OpenPyxlImage(img_out)
            ws_reg.add_image(xl_img, f"E{row_idx}")

    for col in ws_reg.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws_reg.column_dimensions[col_letter].width = min(
            max(max_len + 3, 12), 40
        )

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


with tab3:
    st.subheader("Export Excel Workbook")
    st.caption("Compile stored entries into formatted monthly reports.")

    month_select = st.text_input("Report Month & Year", value="August 2026")

    if st.button(
        "📊 Compile Report File", type="primary", use_container_width=True
    ):
        all_logs = get_all_entries()
        if not all_logs:
            st.warning("No records stored in the database to export.")
        else:
            excel_data = build_excel_report(month_select, all_logs)
            st.download_button(
                label="⬇️ Download .XLSX Report File",
                data=excel_data,
                file_name=f"Monthly_QC_Report_{month_select.replace(' ', '_')}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                use_container_width=True,
            )
